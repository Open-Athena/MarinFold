"""Revise the supplied 36 x 24 inch MoML poster using cached exp325 results.

The original PowerPoint supplies only its logo, QR code, and illustrative
artwork. All poster text remains editable. Numerical panels are redrawn from
the saved summaries, without model inference, rescoring, or new bootstraps.
The output manifest records each displayed summary row and all input hashes.
Run with: uv run --with python-pptx python revise_moml_poster.py --input FILE
"""

import argparse
import hashlib
from io import BytesIO
import json
from pathlib import Path
import subprocess

import matplotlib.pyplot as plt
from matplotlib.colors import LinearSegmentedColormap
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from theme import MARINFOLD_COLOR, METHODS, ORDER

HERE = Path(__file__).resolve().parent
DATA = HERE / "data"
FIGURES = HERE / "plots" / "moml"
TOP7 = HERE.parent / "exp277_models_single_mpnn_pilot/data/eval_rollout_v2"
FONT = "Liberation Sans"
INK = "24333F"
BLACK = "172027"
GRAY = "65717B"
MAROON = MARINFOLD_COLOR.lstrip("#")
BLUE = "4675B9"
TIERS = ["<10", "10–99", "100–999", "≥1000"]
TIER_LABELS = ["<10", "10–100", "100–1000", "≥1000"]
FIGURE_WIDTH = 11.0
PANEL_CAPTIONS = {
    "contact_neff_l2": "314 natural FoldBench proteins; exp277 step266344. Three-way-ranked structure contacts, MarinFold consensus and AFDB Sequence KNN. Equal protein means and 95% bootstrap CIs, regrouped by Neff80. Cached source: data/poster_product_summary.csv; exact displayed rows in data/moml_poster_figure_rows.csv.",
    "structure_neff_l2": "305 common natural FoldBench proteins. Helico with top-L/2 MarinFold or randomly sampled ground truth contacts; other pairs unknown, no MSA. Ground truth averages two subsets. Means and 95% protein-bootstrap CIs from data/neff_summary.csv; exact displayed rows in data/moml_poster_figure_rows.csv.",
    "sampling_l2": "314 matched natural proteins, three MarinFold summaries and seven standard predictors. Mean single, consensus and oracle best individual at P@L/2, plus structure contacts ranked by pLDDT product times degree. All means and 95% CIs copied from data/pl2_sampling_comparison_summary.csv.",
    "top7_contacts": "Top7 (PDB 1QYS), a separate designed-protein illustration outside FoldBench. Experimental contacts and exp277 step266344 contact votes from 100 saved rollouts; source: exp277 data/eval_rollout_v2/top7_provenance.json. Not an independent held-out baseline comparison.",
}


def sha256(path: Path) -> str:
    """Hash a file for the reproducibility manifest."""
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save_panel(fig: plt.Figure, name: str) -> None:
    """Export uncropped, fixed-size artwork for predictable slide placement."""
    for ext in ["png", "pdf", "svg"]:
        metadata = {"Date": None} if ext == "svg" else None
        if ext == "pdf":
            metadata = {"CreationDate": None, "ModDate": None}
        fig.savefig(FIGURES / f"{name}.{ext}", dpi=400, facecolor="white", metadata=metadata)
    (FIGURES / f"{name}.png.meta.json").write_text(json.dumps({
        "script": "revise_moml_poster.py", "args": ["--panels-only"], "caption": PANEL_CAPTIONS[name],
    }, indent=2) + "\n")
    plt.close(fig)


def depth_panel(table: pd.DataFrame, figure: str, metric: str, name: str) -> pd.DataFrame:
    """Draw the saved means and intervals with larger poster typography."""
    methods = ORDER[figure]
    rows = table[(table.figure == figure) & (table.metric == metric)
                 & (table.cohort == "natural") & table.tier.isin(TIERS)
                 & table.method.isin(methods)].copy()
    expected_n = [314, 305][metric == "gdt_ts"]
    assert rows.groupby("method").n.sum().eq(expected_n).all()
    assert set(rows.method) == set(methods)
    fig, ax = plt.subplots(figsize=(FIGURE_WIDTH, 5.45))
    fig.subplots_adjust(left=.115, right=.995, top=.94, bottom=.37)
    for j, method in enumerate(methods):
        r = rows[rows.method == method].set_index("tier").loc[TIERS]
        label, color, marker = METHODS[method]
        x = np.arange(4) + np.linspace(-.29, .29, len(methods))[j]
        ax.errorbar(x, r["mean"], yerr=[r["mean"] - r.ci_low, r.ci_high - r["mean"]],
                    color=color, marker=marker, markersize=6.2, linestyle="none",
                    elinewidth=1.2, capsize=2.3, label=label)
    counts = rows[rows.method == methods[0]].set_index("tier").loc[TIERS, "n"]
    ax.set_xticks(range(4), [f"{tier}\nn = {n}" for tier, n in zip(TIER_LABELS, counts, strict=True)], fontsize=16)
    ax.set(xlim=(-.5, 3.5), ylim=(0, 1.02))
    ax.set_yticks(np.linspace(0, 1, 6))
    ax.set_ylabel("GDT-TS (0–1)" if metric == "gdt_ts" else "Precision @ L/2 contacts", fontsize=18)
    ax.set_xlabel("Effective MSA depth (Neff, 80% identity)", fontsize=17, labelpad=8)
    ax.grid(axis="y", color="#dddddd", linewidth=.7)
    ax.set_axisbelow(True)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.spines["bottom"].set_color("#bbbbbb")
    ax.tick_params(length=0, labelsize=16, pad=7)
    handles, labels = ax.get_legend_handles_labels()
    legend = fig.legend(handles, labels, loc="lower left", bbox_to_anchor=(.09, .005),
               ncol=2, frameon=False, fontsize=16, columnspacing=1.7,
               handlelength=1.3, handletextpad=.5, labelspacing=.45)
    # Leave a measured gap between the x-axis label and the five legend rows.
    fig.canvas.draw()
    renderer = fig.canvas.get_renderer()
    gap_px = 12 * fig.dpi / 72
    needed = legend.get_window_extent(renderer).y1 + gap_px - ax.xaxis.label.get_window_extent(renderer).y0
    if needed > 0:
        pos = ax.get_position()
        shift = needed / fig.bbox.height
        ax.set_position([pos.x0, pos.y0 + shift, pos.width, pos.height - shift])
    fig.canvas.draw()
    assert ax.xaxis.label.get_window_extent().y0 >= legend.get_window_extent().y1 + gap_px - 1
    save_panel(fig, name)
    return rows


def sampling_panel(table: pd.DataFrame) -> pd.DataFrame:
    """Keep all ten methods in the requested bar chart and its original CIs."""
    methods = ["single", "consensus", "best100", "af3", "af2", "boltz2",
               "protenix_msa", "esmfold2", "esmfold", "protenix_ss"]
    rows = table[(table.metric == "p_at_l2") & (table.cohort == "natural")
                 & (table.tier == "All depths")].set_index("method").loc[methods]
    assert rows.n.eq(314).all()
    positions = np.array([0, 1, 2, 3.65, 4.65, 5.65, 6.65, 7.65, 8.65, 9.65])
    labels = ["Mean\nsingle", "Consensus\nof 100", "Oracle\nbest of 100", "AF3\n+ MSA",
              "AF2\n+ MSA", "Boltz-2\n+ MSA", "Protenix-v2\n+ MSA", "ESMFold2", "ESMFold", "Protenix-v2\nsingle seq."]
    colors = [METHODS[m][1] for m in methods]
    fig, ax = plt.subplots(figsize=(FIGURE_WIDTH, 3.55))
    fig.subplots_adjust(left=.115, right=.995, bottom=.31, top=.90)
    ax.bar(positions, rows["mean"], width=.72, color=colors, zorder=3)
    ax.errorbar(positions, rows["mean"], yerr=[rows["mean"] - rows.ci_low, rows.ci_high - rows["mean"]],
                fmt="none", ecolor="#333333", capsize=3, elinewidth=1.2, zorder=4)
    for x, (_, row), color in zip(positions, rows.iterrows(), colors, strict=True):
        ax.text(x, row.ci_high + .025, f"{row['mean']:.3f}", ha="center", fontsize=13.5, color=color)
    ax.set_xticks(positions, labels, fontsize=12)
    ax.set(xlim=(-.65, 10.3), ylim=(0, 1.05))
    ax.set_yticks([0, .2, .4, .6, .8, 1])
    ax.set_ylabel("Precision @ L/2 contacts", fontsize=18)
    ax.axvline(2.85, color="#bbbbbb", linestyle="--", linewidth=.8)
    ax.grid(axis="y", color="#dddddd", zorder=0)
    ax.spines[["top", "right", "left"]].set_visible(False)
    ax.spines["bottom"].set_color("#bbbbbb")
    ax.tick_params(length=0, pad=7)
    ax.text(1, -.41, "MarinFold · 100 rollouts", transform=ax.get_xaxis_transform(),
            ha="center", fontsize=16, color=MARINFOLD_COLOR)
    ax.text(6.65, -.41, "Standard predictors", transform=ax.get_xaxis_transform(),
            ha="center", fontsize=16, color="#24333F")
    save_panel(fig, "sampling_l2")
    return rows.reset_index()


def top7_panel() -> None:
    """Redraw the demonstration from the evaluated checkpoint's cached votes."""
    truth = json.loads((TOP7 / "top7_ground_truth.json").read_text())
    provenance = json.loads((TOP7 / "top7_provenance.json").read_text())
    assert provenance["step"] == 266344 and truth["L"] == 92
    votes = np.loadtxt(TOP7 / "top7_votes.csv", delimiter=",")
    matrix = np.zeros_like(votes)
    for i, j, degree in truth["contacts"]:
        if degree >= .001 and abs(i-j) >= 6:
            matrix[i, j] = matrix[j, i] = 1
    fig, axes = plt.subplots(1, 2, figsize=(8.0, 3.3))
    fig.subplots_adjust(left=.07, right=.85, top=.85, bottom=.18, wspace=.25)
    axes[0].imshow(matrix, origin="lower", cmap="Greys", vmin=0, vmax=1)
    cmap = LinearSegmentedColormap.from_list("contact_votes", ["#ffffff", "#e4b981", MARINFOLD_COLOR, "#1b0f16"])
    im = axes[1].imshow(votes / 100, origin="lower", cmap=cmap, vmin=0, vmax=1)
    for ax, title in zip(axes, ["Ground truth", "MarinFold consensus"], strict=True):
        ax.set_title(title, fontsize=17)
        ax.set_xticks([0, 46, 91])
        ax.set_yticks([0, 46, 91])
        ax.set_xlabel("Residue index", fontsize=15)
    axes[0].set_ylabel("Residue index", fontsize=15)
    cax = fig.add_axes([.86, .18, .018, .67])
    fig.colorbar(im, cax=cax, ticks=[0, .5, 1])
    cax.set_ylabel("Fraction of 100 rollouts", fontsize=14)
    save_panel(fig, "top7_contacts")


def rgb(value: str) -> RGBColor:
    """Convert a hex palette entry into a native PowerPoint color."""
    return RGBColor.from_string(value)


def text(slide, x: float, y: float, w: float, h: float, content: str,
         size: float = 20, color: str = INK, bold: bool = False,
         align=PP_ALIGN.LEFT, gap: float = 6):
    """Add editable, explicitly formatted text without font substitution."""
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, line in enumerate(content.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.space_after = Pt(gap)
        p.line_spacing = 1.08
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(color)
        p.font.bold = bold
    return shape


def heading(slide, x: float, y: float, content: str, w: float = 11, h: float = .53) -> None:
    """Use a consistent hierarchy for each poster section."""
    text(slide, x, y, w, h, content, 28, BLACK, True, gap=0)


def asset(slide, source, x: float, y: float, w: float, h: float,
          crop: tuple[float, float, float, float] | None = None) -> None:
    """Reuse the user's artwork through PowerPoint crop fields, not pixel edits."""
    shape = slide.shapes.add_picture(BytesIO(source.image.blob), Inches(x), Inches(y), Inches(w), Inches(h))
    shape.crop_left, shape.crop_top, shape.crop_right, shape.crop_bottom = (
        crop if crop is not None else (source.crop_left, source.crop_top, source.crop_right, source.crop_bottom))


def box(slide, x: float, y: float, w: float, h: float, label: str,
        size: float = 18, color: str = BLUE) -> None:
    """Draw an editable architecture component."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb("FFFFFF")
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(1.2)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(.06)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, line in enumerate(label.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(BLACK)


def arrow(slide, x: float, y: float, w: float, h: float, down: bool = False,
          left: bool = False, up: bool = False) -> None:
    """Add a native editable arrow in the original diagram's blue."""
    kind = MSO_SHAPE.DOWN_ARROW if down else (MSO_SHAPE.LEFT_ARROW if left else (MSO_SHAPE.UP_ARROW if up else MSO_SHAPE.RIGHT_ARROW))
    shape = slide.shapes.add_shape(kind,
                                  Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(BLUE)
    shape.line.fill.background()


def training_table(slide) -> None:
    """Use the audited epoch counts and label redesigns as examples."""
    values = [
        ["Training source", "Examples (M)"],
        ["2026 ESM Atlas · ESMFold2 [1]", "65.55"],
        ["AlphaFold DB · AlphaFold2 [2]", "3.96"],
        ["ProteinMPNN redesigns [3]\nof retained parent backbones", "162.57"],
        ["Total (one epoch)", "232.09"],
    ]
    s = slide.shapes.add_table(len(values), 2, Inches(1.08), Inches(15.36), Inches(6.25), Inches(2.38))
    s.table.columns[0].width = Inches(4.57)
    s.table.columns[1].width = Inches(1.68)
    for i, row in enumerate(values):
        s.table.rows[i].height = Inches(.63 if i == 3 else .4375)
        for j, label in enumerate(row):
            c = s.table.cell(i, j)
            c.text = label
            c.fill.solid()
            c.fill.fore_color.rgb = rgb("F2F2F2" if i in [0, 4] else "FFFFFF")
            c.margin_left = c.margin_right = Inches(.08)
            c.margin_top = c.margin_bottom = Inches(.045)
            for p in c.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(16.5)
                p.font.bold = i in [0, 4]
                p.font.color.rgb = rgb(INK)
                p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.RIGHT
    text(slide, 1.08, 17.85, 6.3, .52,
         "Counts after filtering; rounded independently. Redesigns add sequences, not new backbones.",
         15.5, gap=0)


REFERENCES = [
    ("[1] Candido et al. (2026). ESMC, ESMFold2 and ESM Atlas. bioRxiv. doi:10.64898/2026.06.03.729735",
     "https://doi.org/10.64898/2026.06.03.729735"),
    ("[2] Varadi et al. (2024). AlphaFold DB. Nucleic Acids Res. 52:D368–375. doi:10.1093/nar/gkad1011",
     "https://doi.org/10.1093/nar/gkad1011"),
    ("[3] Dauparas et al. (2022). ProteinMPNN. Science 378:49–56. doi:10.1126/science.add2187",
     "https://doi.org/10.1126/science.add2187"),
    ("[4] Zhang et al. (2026). Protenix-v1. bioRxiv. doi:10.64898/2026.02.05.703733",
     "https://doi.org/10.64898/2026.02.05.703733"),
    ("[5] Xu et al. (2026). FoldBench. Nat. Commun. 17:442. doi:10.1038/s41467-025-67127-3",
     "https://doi.org/10.1038/s41467-025-67127-3"),
    ("[6] Jumper et al. (2021). AlphaFold2. Nature 596:583–589. doi:10.1038/s41586-021-03819-2",
     "https://doi.org/10.1038/s41586-021-03819-2"),
    ("[7] Abramson et al. (2024). AlphaFold3. Nature 630:493–500. doi:10.1038/s41586-024-07487-w",
     "https://doi.org/10.1038/s41586-024-07487-w"),
    ("[8] Passaro et al. (2025). Boltz-2. bioRxiv. doi:10.1101/2025.06.14.659707",
     "https://doi.org/10.1101/2025.06.14.659707"),
    ("[9] Lin et al. (2023). ESMFold. Science 379:1123–1130. doi:10.1126/science.ade2574",
     "https://doi.org/10.1126/science.ade2574"),
    ("[10] Zhang et al. (2026). Protenix-v2. bioRxiv. doi:10.64898/2026.04.10.717613",
     "https://doi.org/10.64898/2026.04.10.717613"),
    ("[11] Yang et al. (2025). Qwen3 Technical Report. arXiv:2505.09388.",
     "https://arxiv.org/abs/2505.09388"),
    ("[12] Hall et al. (2025). Introducing Marin: An Open Lab for Building Foundation Models. Marin Community.",
     "https://marin.community/blog/2025/05/19/announcement/"),
]


def build_poster(original: Path, output: Path) -> None:
    """Keep the poster's three-column organization while replacing its copy."""
    old = Presentation(original)
    assert len(old.slides) == 1 and old.slide_width == Inches(36) and old.slide_height == Inches(24)
    original_shapes = {s.shape_id: s for s in old.slides[0].shapes}
    prs = Presentation()
    prs.slide_width, prs.slide_height = old.slide_width, old.slide_height
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb("FFFFFF")
    text(slide, 1.08, .82, 16.0, 2.15,
         "Protein structure prediction by autoregressive modeling of residue contacts", 44, BLACK, gap=0)
    text(slide, 17.72, .92, 13.3, 2.05,
         "Timothy O’Donnell¹, Eric Czech¹, Zack Nichols¹, Jacob Silterra², James Roney³, Ashley Gin³, Sergey Ovchinnikov³\n"
         "¹ Open Athena    ² TetraScience    ³ MIT\nContact: tim@openathena.ai", 27, BLACK, gap=6)
    asset(slide, original_shapes[3], 31.45, .38, 4.0, 1.637)

    # Left column: shorten the motivation and spend the space on the protocol.
    heading(slide, 1.08, 3.82, "Summary")
    text(slide, 1.08, 4.57, 11.05, 4.28,
         "MarinFold predicts residue–residue contacts from a single sequence, without MSAs or pretrained protein-language-model embeddings at inference. Training still inherits information from sequences and teacher models.\n"
         "Consensus across 100 samples raises mean precision @ L/2 from 0.414 to 0.714 on 314 natural FoldBench proteins.\n"
         "Accuracy remains associated with effective MSA depth (Neff), despite omitting MSAs.\n"
         "Helico reaches mean GDT-TS 0.478 with MarinFold’s L/2 contacts versus 0.843 with ground truth contacts, suggesting that contact quality limits accuracy.",
         22, gap=12)
    heading(slide, 1.08, 9.00, "Evaluation and training exclusion")
    text(slide, 1.08, 9.72, 11.05, 3.88,
         "Cohort. FoldBench’s 334 monomers [5] minus 19 designs and one context-overflow target (8uxt_A) gives 314 natural proteins: 97 validation + 217 test. Structure plots use 305 shared complete cases; nine lack a complete set of archived results.\n"
         "Exclusion. Native training sequences matching evaluation references at ≥30% identity over ≥50% of the shorter sequence are removed before redesign. No structural-similarity filter is applied. Helico’s PDB training cutoff is 2021-09-30; teacher-model training overlap is not controlled.\n"
         "Neff80 = Σᵢ 1/nᵢ: nᵢ counts ≥80%-identical alignment sequences, including itself (identity over shared non-gap columns). Means weight proteins equally; error bars are 95% protein-bootstrap intervals (5,000 resamples). These exploratory analyses reuse validation and test data.\n"
         "Sequence KNN transfers and votes on contacts from 10 aligned nearest native AFDB training sequences after exclusion; it does not search the redesign corpus.",
         18, gap=8)
    heading(slide, 1.08, 14.06, "Distilling predicted structures", h=.6)
    text(slide, 1.08, 14.78, 11.05, .44,
         "AlphaFold2 and ESMFold2 supply the structural training labels.", 20)
    training_table(slide)
    asset(slide, original_shapes[97], 9.02, 15.67, 1.64, 1.69)
    arrow(slide, 7.57, 16.30, 1.16, .18)
    text(slide, 8.4, 17.46, 2.9, .37, "Contact maps", 18.5, bold=True, align=PP_ALIGN.CENTER)
    arrow(slide, 9.72, 17.96, .18, .42, down=True)
    asset(slide, original_shapes[9], 7.85, 18.62, 3.85, 2.02)
    text(slide, 7.78, 20.65, 4.12, .46, "Each bracketed symbol is one token.", 15.5, align=PP_ALIGN.CENTER)
    arrow(slide, 6.71, 19.46, .83, .18, left=True)
    text(slide, 1.08, 18.61, 5.52, 2.35,
         "MarinFold · 1.47B parameters\nDecoder-only Qwen3 architecture [11]\nTrained from scratch with Marin [12]\nContext length: 8,192 tokens\n232.09M documents · 248.58B raw tokens\nOne epoch · checkpoint step 266,344",
         19, gap=5)

    # Middle column: preserve a separate Top7 illustration, then the two results.
    heading(slide, 12.94, 3.82, "From a sequence to a contact map")
    asset(slide, original_shapes[9], 13.03, 5.34, 2.34, 1.77,
          crop=(.324, .281, .516, .176))
    slide.shapes.add_picture(str(FIGURES / "top7_contacts.png"), Inches(15.31), Inches(4.93), Inches(8.0), Inches(3.3))
    text(slide, 12.98, 8.35, 11.0, .66,
         "Top7 (PDB: 1QYS), a computationally designed protein. Separate illustration outside FoldBench; shown with the same 248B-token checkpoint.",
         18, gap=0)
    heading(slide, 12.94, 9.27, "Consensus improves contact precision")
    slide.shapes.add_picture(str(FIGURES / "sampling_l2.png"), Inches(12.81), Inches(9.93), Inches(11.0), Inches(3.55))
    text(slide, 12.98, 13.53, 11.0, 1.57,
         "P@L/2: k = floor(L/2); L is input length. Score resolved pairs ≥6 residues apart against ConFind truth (3 Å; degree ≥0.001). Baselines rank positive ConFind pairs by pLDDTᵢ × pLDDTⱼ × degree; unfilled ranks score zero.\n"
         "Consensus votes across 100 rollouts; oracle selects the best individual using truth. Sampling consensus includes unfinished maps; depth consensus excludes them.",
         17, gap=6)
    heading(slide, 12.94, 15.37, "Accuracy tracks effective MSA depth")
    slide.shapes.add_picture(str(FIGURES / "contact_neff_l2.png"), Inches(12.81), Inches(16.02), Inches(11.0), Inches(5.45))

    # Right column: an editable diagram makes the MSA-free input explicit.
    heading(slide, 24.59, 3.82, "Generating 3D structures with Helico")
    text(slide, 24.59, 4.58, 10.9, 1.01,
         "We fine-tune Protenix-v1 [4] to condition on three-state contact maps without an MSA. The original confidence head is unchanged.",
         21, gap=0)
    box(slide, 24.64, 6.03, 2.77, .85, "Sequence, conformers,\nbonds · no MSA", 18)
    box(slide, 28.0, 6.03, 1.53, .85, "Input\nembedder", 18)
    box(slide, 31.63, 6.03, 1.77, .85, "Pairformer\nstack", 18)
    box(slide, 33.98, 6.03, 1.56, .85, "Diffusion\nmodule", 18)
    arrow(slide, 27.46, 6.35, .47, .16)
    arrow(slide, 29.59, 6.35, 1.95, .16)
    text(slide, 29.68, 5.95, 1.83, .30, "Single features", 14.5, align=PP_ALIGN.CENTER)
    arrow(slide, 33.46, 6.35, .45, .16)
    asset(slide, original_shapes[164], 24.85, 7.39, 1.14, 1.17)
    text(slide, 24.59, 8.62, 3.3, .52, "Contact map\npresent / absent / unknown", 15.5, bold=True, align=PP_ALIGN.CENTER, gap=0)
    box(slide, 26.64, 7.64, 1.63, .82, "Linear\n3 → 128", 18, MAROON)
    arrow(slide, 26.08, 7.96, .47, .16)
    arrow(slide, 28.32, 7.96, .35, .16)
    text(slide, 28.70, 7.71, .53, .51, "+", 34, BLACK, True, align=PP_ALIGN.CENTER)
    arrow(slide, 28.86, 6.97, .16, .61, down=True)
    text(slide, 29.09, 7.02, 1.28, .55, "Pair features", 14.5, gap=0)
    arrow(slide, 29.30, 7.96, 3.27, .13)
    arrow(slide, 32.44, 6.97, .13, 1.11, up=True)
    text(slide, 29.65, 8.24, 3.0, .42, "Contact-conditioned pair features", 15, align=PP_ALIGN.CENTER)
    heading(slide, 24.59, 9.31, "Structure prediction accuracy")
    slide.shapes.add_picture(str(FIGURES / "structure_neff_l2.png"), Inches(24.46), Inches(9.97), Inches(11.0), Inches(5.45))
    text(slide, 24.59, 15.55, 10.88, 1.64,
         "Helico receives top-L/2 MarinFold contacts or L/2 uniformly sampled ground truth contacts; all other pairs stay unknown. Three diffusion samples per map are selected by the original ranking_score. Ground truth uses two subsets, averaged within protein.\n"
         "ESMFold and ESMFold2 use no MSA. +MSA arms share frozen alignments; AF2/AF3/Boltz-2 disable templates. Inference budgets differ.",
         17.5, gap=7)
    heading(slide, 24.59, 17.43, "Conclusions and next steps")
    text(slide, 24.59, 18.14, 10.88, 2.84,
         "Ground truth contacts support high accuracy even at Neff <10, suggesting that predicted-contact quality is a major bottleneck.\n"
         "MarinFold + Helico exceeds single-sequence Protenix-v2 mainly at higher Neff. At Neff <10, means are 0.285 versus 0.290. Stronger baselines remain ahead.\n"
         "Helico without contacts is shown. A matched pretrained Protenix-v1 baseline is still needed to isolate fine-tuning from contact conditioning.\n"
         "Next: useful sample diversity, training-data composition, inference-time search, post-training and protein–protein complexes.",
         19.5, gap=8)

    # Compact real citations replace the template references and off-slide notes.
    heading(slide, 1.08, 21.62, "References", w=5, h=.40)
    for i, (label, url) in enumerate(REFERENCES):
        col, row = divmod(i, 4)
        s = text(slide, 1.08 + col * 10.60, 22.14 + row * .39, 10.28, .37, label, 12.5, gap=0)
        s.text_frame.paragraphs[0].runs[0].hyperlink.address = url
        s.text_frame.paragraphs[0].runs[0].font.color.rgb = rgb(INK)
        s.text_frame.paragraphs[0].runs[0].font.underline = False
    asset(slide, original_shapes[16], 33.18, 21.86, 1.96, 1.96)
    text(slide, 27.35, 21.61, 5.42, .38, "Openly developed · contributions welcome", 15, color=MAROON, align=PP_ALIGN.RIGHT)
    link = text(slide, 22.30, 23.74, 10.4, .22, "Code, methods and full references: github.com/Open-Athena/MarinFold", 12, gap=0)
    link.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://github.com/Open-Athena/MarinFold/tree/exp/325-writeup-analysis/experiments/exp325_evals_writeup_analysis"
    slide.notes_slide.notes_text_frame.text = (
        "Revised from the user-supplied MarinFold MoML poster. Numerical results are cached exp325 "
        "summaries; no new inference, scores or confidence intervals. See MOML_POSTER_REVIEW.md for "
        "each review decision, exact exclusions, full citations and limitations. Figures: "
        "revise_moml_poster.py --input <original poster.pptx>. The Top7 illustration reuses exp277 "
        "step266344 votes. Original source artwork is preserved only for the logo, QR, molecular "
        "cartoon, illustrative token document and schematic contact map."
    )
    prs.core_properties.title = "MarinFold · MoML poster (revised)"
    prs.core_properties.subject = "Sequence-only contact prediction and contact-conditioned folding"
    prs.core_properties.author = old.core_properties.author
    for s in slide.shapes:
        assert 0 <= s.left and 0 <= s.top
        assert s.left + s.width <= prs.slide_width and s.top + s.height <= prs.slide_height, s.name
    prs.save(output)


def main() -> None:
    """Produce an editable PowerPoint, print PDF, preview and source manifest."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path)
    parser.add_argument("--panels-only", action="store_true", help="Restyle the cached figures without the supplied PowerPoint")
    parser.add_argument("--output-dir", type=Path, default=HERE / "scratch/moml_poster/revised")
    args = parser.parse_args()
    if args.input is None and not args.panels_only:
        parser.error("--input is required unless --panels-only is supplied")
    training = pd.read_csv(DATA / "training_sources.csv")
    assert training.documents.sum() == 232090905
    assert training.raw_tokens.sum() == 248583762834
    args.output_dir.mkdir(parents=True, exist_ok=True)
    FIGURES.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update({"font.family": FONT, "font.size": 15, "pdf.fonttype": 42,
                         "svg.fonttype": "none", "text.color": "#24333F", "axes.labelcolor": "#24333F"})
    used = []
    for filename, figure, metric, name in [
        ("poster_product_summary.csv", "04_contacts_pl2_neff", "p_at_l2", "contact_neff_l2"),
        ("neff_summary.csv", "05_folding_neff", "gdt_ts", "structure_neff_l2"),
    ]:
        table = pd.read_csv(DATA / filename)
        table["source_summary_row"] = table.index
        rows = depth_panel(table, figure, metric, name)
        rows["source_summary"] = filename
        rows["poster_panel"] = name
        used.append(rows)
    filename = "pl2_sampling_comparison_summary.csv"
    table = pd.read_csv(DATA / filename)
    table["source_summary_row"] = table.index
    rows = sampling_panel(table)
    rows["source_summary"] = filename
    rows["poster_panel"] = "sampling_l2"
    used.append(rows)
    top7_panel()
    pd.concat(used, ignore_index=True).to_csv(DATA / "moml_poster_figure_rows.csv", index=False)
    if args.panels_only:
        print(FIGURES)
        return
    output = args.output_dir / "MarinFold_MoML_revised.pptx"
    build_poster(args.input, output)
    subprocess.run(["libreoffice", "-env:UserInstallation=file:///tmp/marinfold-moml-revised",
                    "--headless", "--convert-to", "pdf", "--outdir", str(args.output_dir), str(output)], check=True)
    pdf = output.with_suffix(".pdf")
    subprocess.run(["pdftoppm", "-scale-to", "3000", "-singlefile", "-png", str(pdf),
                    str(args.output_dir / "MarinFold_MoML_revised_preview")], check=True)
    subprocess.run(["pdftoppm", "-r", "300", "-W", "10800", "-H", "7200", "-singlefile", "-png", str(pdf),
                    str(args.output_dir / "MarinFold_MoML_revised")], check=True)
    sources = [DATA / name for name in ["training_sources.csv", "poster_product_summary.csv", "neff_summary.csv",
               "pl2_sampling_comparison_summary.csv", "structure_exclusions.csv"]]
    sources.extend(TOP7 / name for name in ["top7_votes.csv", "top7_ground_truth.json", "top7_provenance.json"])
    manifest = {"original_pptx_sha256": sha256(args.input), "checkpoint": "exp277-step-266344",
                "raw_training_tokens": 248583762834, "training_documents": 232090905,
                "dimensions_inches": [36, 24], "numerical_reanalysis": False,
                "inputs": {str(p.relative_to(HERE.parent.parent)): sha256(p) for p in sources},
                "figure_rows": "data/moml_poster_figure_rows.csv",
                "outputs": {p.name: sha256(p) for p in args.output_dir.glob("MarinFold_MoML_revised*.*")},
                "generator": "revise_moml_poster.py --input <original poster.pptx>"}
    (DATA / "moml_poster_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    print(output)
    print(pdf)


if __name__ == "__main__":
    main()
